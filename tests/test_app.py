"""Integration tests to run after installing requirements: python -m unittest discover -s tests -v"""
import io
import os
import unittest

os.environ.pop("RENDER", None)
os.environ.pop("ADMIN_PASSWORD", None)
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["SECRET_KEY"] = "test-secret-key-only"

from openpyxl import load_workbook
from app import app, db


class TrackerTests(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.ctx = app.app_context()
        self.ctx.push()
        db.drop_all()
        db.create_all()
        self.client = app.test_client()
        self.client.get("/")
        with self.client.session_transaction() as session:
            self.csrf = session["csrf"]
        self.headers = {"X-CSRF-Token": self.csrf}

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.ctx.pop()

    def post(self, url, data):
        return self.client.post(url, json=data, headers=self.headers)

    def create_jira(self):
        res = self.post("/api/jiras", {"sprint": "Sprint 1", "jira": "OPPC-298", "story_points": 5})
        self.assertEqual(res.status_code, 201)
        return res.get_json()["jira"]["id"]

    def test_conversion_and_duplicate(self):
        self.create_jira()
        row = self.client.get("/api/jiras").get_json()["jiras"][0]
        self.assertEqual((row["days"], row["total_hours"]), (2.5, 20))
        res = self.post("/api/jiras", {"sprint": "Sprint 2", "jira": "OPPC-298", "story_points": 8})
        self.assertEqual(res.status_code, 409)

    def test_bookings_and_running_balance(self):
        id = self.create_jira()
        for when, hours in [("2026-10-10", 4), ("2026-10-11", 6)]:
            res = self.post(f"/api/jiras/{id}/bookings", {"date": when, "hours": hours, "remark": "Work"})
            self.assertEqual(res.status_code, 201)
        data = self.client.get(f"/api/jiras/{id}").get_json()["jira"]
        self.assertEqual([r["remaining_hours"] for r in data["bookings"]], [16.0, 10.0])
        self.assertEqual(data["remaining_hours"], 10.0)
        bad = self.post(f"/api/jiras/{id}/bookings", {"date": "bad", "hours": -1})
        self.assertEqual(bad.status_code, 400)

    def test_excel_export_contains_jira_sheet_and_values(self):
        id = self.create_jira()
        self.post(f"/api/jiras/{id}/bookings", {"date": "2026-10-10", "hours": 4, "remark": "Development"})
        response = self.client.get("/export/excel")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb.sheetnames, ["Master Tracker", "OPPC-298"])
        self.assertEqual(wb["Master Tracker"]["E5"].value, 20.0)
        self.assertEqual(wb["OPPC-298"]["E5"].value, 16.0)
        self.assertEqual(wb["OPPC-298"]["D5"].value, "Development")

    def test_csrf_is_required_for_mutations(self):
        response = self.client.post("/api/jiras", json={"sprint":"S1","jira":"ABC-123","story_points":2})
        self.assertEqual(response.status_code, 403)


if __name__ == "__main__":
    unittest.main()
