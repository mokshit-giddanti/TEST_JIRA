"""JIRA & BCS Hours Tracker: Flask web app with persistent SQL storage."""
import hmac
import io
import os
import re
import secrets
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from uuid import uuid4

from flask import Flask, abort, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import selectinload

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = bool(os.environ.get("RENDER"))
app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024
_database_url = os.environ.get("DATABASE_URL", "sqlite:///jira_bcs_local.db")
if _database_url.startswith("postgres://"):
    _database_url = _database_url.replace("postgres://", "postgresql://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = _database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
db = SQLAlchemy(app)
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
if os.environ.get("RENDER") and not ADMIN_PASSWORD:
    raise RuntimeError("ADMIN_PASSWORD must be set for a Render deployment.")


class Jira(db.Model):
    __tablename__ = "jiras"
    id = db.Column(db.String(36), primary_key=True, default=lambda: str(uuid4()))
    sprint = db.Column(db.String(64), nullable=False)
    code = db.Column(db.String(50), nullable=False, unique=True, index=True)
    story_points = db.Column(db.Numeric(9, 2), nullable=False)
    created_at = db.Column(db.DateTime, nullable=False, default=lambda: datetime.now(timezone.utc).replace(tzinfo=None))
    bookings = db.relationship("Booking", back_populates="jira", cascade="all, delete-orphan")


class Booking(db.Model):
    __tablename__ = "bookings"
    id = db.Column(db.String(36), primary_key=True, default=lambda: str(uuid4()))
    jira_id = db.Column(db.String(36), db.ForeignKey("jiras.id", ondelete="CASCADE"), nullable=False, index=True)
    book_date = db.Column(db.Date, nullable=False)
    hours = db.Column(db.Numeric(9, 2), nullable=False)
    remark = db.Column(db.String(500), nullable=False, default="")
    created_at = db.Column(db.DateTime, nullable=False, default=lambda: datetime.now(timezone.utc).replace(tzinfo=None))
    jira = db.relationship("Jira", back_populates="bookings")


def money(value):
    return float(value.quantize(Decimal("0.01")))


def sorted_bookings(jira):
    return sorted(jira.bookings, key=lambda b: (b.book_date, b.created_at, b.id))


def jira_dict(jira, include_bookings=False):
    points = jira.story_points
    days = points / 2
    total = points * 4
    ordered = sorted_bookings(jira)
    booked = sum((b.hours for b in ordered), Decimal(0))
    result = {
        "id": jira.id,
        "sprint": jira.sprint,
        "jira": jira.code,
        "story_points": money(points),
        "days": money(days),
        "total_hours": money(total),
        "booked_hours": money(booked),
        "remaining_hours": money(total - booked),
        "booking_count": len(ordered),
    }
    if include_bookings:
        running = total
        result["bookings"] = []
        for b in ordered:
            running -= b.hours
            result["bookings"].append({
                "id": b.id,
                "date": b.book_date.isoformat(),
                "hours": money(b.hours),
                "remark": b.remark,
                "remaining_hours": money(running),
            })
    return result


def json_error(message, status=400):
    return jsonify({"error": message}), status


def valid_decimal(value, name, *, min_value=Decimal("0"), max_value=Decimal("10000"), allow_zero=False):
    try:
        val = Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"{name} must be a valid number.")
    if not val.is_finite() or val > max_value or val < min_value or (val == 0 and not allow_zero):
        raise ValueError(f"{name} must be greater than 0 and at most {max_value}.")
    return val


def validated_jira(data):
    sprint = str(data.get("sprint") or "").strip()
    code = str(data.get("jira") or "").strip().upper()
    if not sprint or len(sprint) > 64:
        raise ValueError("Sprint is required (maximum 64 characters).")
    if not re.fullmatch(r"[A-Z0-9][A-Z0-9._-]{1,49}", code):
        raise ValueError("Enter a valid JIRA number, such as OPPC-298.")
    points = valid_decimal(data.get("story_points"), "Story Points")
    return sprint, code, points


def validated_booking(data):
    try:
        booking_date = date.fromisoformat(str(data.get("date") or ""))
    except ValueError:
        raise ValueError("Enter a valid date.")
    hours = valid_decimal(data.get("hours"), "BCS booking hours")
    remark = str(data.get("remark") or "").strip()
    if len(remark) > 500:
        raise ValueError("Remark must be 500 characters or fewer.")
    return booking_date, hours, remark


@app.before_request
def authenticate_and_protect():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(24)
    if request.path.startswith("/static/") or request.path in ("/healthz", "/login"):
        return
    if ADMIN_PASSWORD and not session.get("authenticated"):
        if request.path.startswith("/api/") or request.path.startswith("/export/"):
            return json_error("Sign in to access the tracker.", 401)
        return redirect(url_for("login"))
    if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        supplied = request.headers.get("X-CSRF-Token", "") or request.form.get("csrf", "")
        if not hmac.compare_digest(supplied, session.get("csrf", "")):
            return json_error("Security token expired. Reload the page and retry.", 403)


@app.get("/healthz")
def healthz():
    return "ok", 200


@app.route("/login", methods=["GET", "POST"])
def login():
    if not ADMIN_PASSWORD:
        return redirect(url_for("home"))
    if request.method == "POST":
        supplied = request.form.get("password", "")
        token = request.form.get("csrf", "")
        if hmac.compare_digest(token, session.get("csrf", "")) and hmac.compare_digest(supplied, ADMIN_PASSWORD):
            session.clear()
            session["csrf"] = secrets.token_hex(24)
            session["authenticated"] = True
            return redirect(url_for("home"))
        return render_template("login.html", error="Incorrect password. Try again.", csrf=session["csrf"]), 401
    return render_template("login.html", error=None, csrf=session["csrf"])


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.get("/")
def home():
    return render_template("index.html", csrf=session["csrf"], auth_enabled=bool(ADMIN_PASSWORD))


@app.get("/api/jiras")
def get_jiras():
    jiras = db.session.execute(select(Jira).options(selectinload(Jira.bookings)).order_by(Jira.created_at.desc(), Jira.code)).scalars().all()
    return jsonify({"jiras": [jira_dict(j) for j in jiras]})


@app.post("/api/jiras")
def add_jira():
    data = request.get_json(silent=True) or {}
    try:
        sprint, code, points = validated_jira(data)
    except ValueError as e:
        return json_error(str(e))
    if db.session.execute(select(Jira.id).where(Jira.code == code)).scalar_one_or_none():
        return json_error("This JIRA already exists.", 409)
    jira = Jira(sprint=sprint, code=code, story_points=points)
    db.session.add(jira)
    db.session.commit()
    return jsonify({"jira": jira_dict(jira)}), 201


@app.get("/api/jiras/<jira_id>")
def get_jira(jira_id):
    jira = db.session.execute(select(Jira).where(Jira.id == jira_id).options(selectinload(Jira.bookings))).scalar_one_or_none()
    if not jira:
        return json_error("JIRA not found.", 404)
    return jsonify({"jira": jira_dict(jira, include_bookings=True)})


@app.put("/api/jiras/<jira_id>")
def update_jira(jira_id):
    jira = db.session.get(Jira, jira_id)
    if not jira:
        return json_error("JIRA not found.", 404)
    try:
        sprint, code, points = validated_jira(request.get_json(silent=True) or {})
    except ValueError as e:
        return json_error(str(e))
    duplicate = db.session.execute(select(Jira.id).where(Jira.code == code, Jira.id != jira_id)).scalar_one_or_none()
    if duplicate:
        return json_error("This JIRA number is already used.", 409)
    jira.sprint, jira.code, jira.story_points = sprint, code, points
    db.session.commit()
    db.session.refresh(jira)
    return jsonify({"jira": jira_dict(jira)})


@app.delete("/api/jiras/<jira_id>")
def delete_jira(jira_id):
    jira = db.session.get(Jira, jira_id)
    if not jira:
        return json_error("JIRA not found.", 404)
    db.session.delete(jira)
    db.session.commit()
    return jsonify({"ok": True})


@app.post("/api/jiras/<jira_id>/bookings")
def add_booking(jira_id):
    jira = db.session.get(Jira, jira_id)
    if not jira:
        return json_error("JIRA not found.", 404)
    try:
        booking_date, hours, remark = validated_booking(request.get_json(silent=True) or {})
    except ValueError as e:
        return json_error(str(e))
    b = Booking(jira_id=jira_id, book_date=booking_date, hours=hours, remark=remark)
    db.session.add(b)
    db.session.commit()
    return jsonify({"booking_id": b.id}), 201


@app.put("/api/bookings/<booking_id>")
def update_booking(booking_id):
    b = db.session.get(Booking, booking_id)
    if not b:
        return json_error("Booking not found.", 404)
    try:
        b.book_date, b.hours, b.remark = validated_booking(request.get_json(silent=True) or {})
    except ValueError as e:
        return json_error(str(e))
    db.session.commit()
    return jsonify({"ok": True})


@app.delete("/api/bookings/<booking_id>")
def delete_booking(booking_id):
    b = db.session.get(Booking, booking_id)
    if not b:
        return json_error("Booking not found.", 404)
    db.session.delete(b)
    db.session.commit()
    return jsonify({"ok": True})


NAVY = "19263A"
BLUE = "2563EB"
PALE = "EEF4FF"
WHITE = "FFFFFF"
GRAY = "64748B"
GREEN = "13795B"
BORDER = "DCE5F0"


def excel_text(value):
    """Prevent entered text from becoming an Excel formula."""
    s = str(value)
    return "'" + s if s.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else s


def style_sheet(ws, headers, widths):
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=4, max_row=4, max_col=len(headers)):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Aptos", color=WHITE, bold=True, size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=BORDER))
    ws.row_dimensions[4].height = 34
    for row in ws.iter_rows(min_row=5, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name="Aptos", color=NAVY, size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=BORDER))
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
        ws.row_dimensions[row[0].row].height = 24
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(5, ws.max_row)}"
    for cell in ws[1]:
        cell.font = Font(name="Aptos Display", size=19, bold=True, color=NAVY)
    for cell in ws[2]:
        cell.font = Font(name="Aptos", size=10, color=GRAY)


def export_workbook(jiras):
    """Downloadable Excel workbook: master list and a sheet per JIRA."""
    wb = Workbook()
    master = wb.active
    master.title = "Master Tracker"
    master.merge_cells("A1:F1")
    master["A1"] = "JIRA & BCS HOURS TRACKER"
    master.merge_cells("A2:F2")
    master["A2"] = "1 Story Point = 0.5 Day = 4 Hours  |  1 Day = 8 Hours"
    headers = ["Sprint", "JIRA", "Story Points", "Days", "Total Hours", "Bookings"]
    master.append([])
    master.append(headers)
    # Excel sheet names max 31 chars; avoid collisions for long JIRA keys.
    used_names = {"Master Tracker"}
    for jira in jiras:
        proposed = re.sub(r"[\\/*?:\[\]]", "_", jira.code)[:31]
        sheet_name = proposed
        suffix = 2
        while sheet_name.lower() in {n.lower() for n in used_names}:
            tail = f"-{suffix}"
            sheet_name = proposed[:31 - len(tail)] + tail
            suffix += 1
        used_names.add(sheet_name)
        mrow = master.max_row + 1
        master.append([excel_text(jira.sprint), jira.code, money(jira.story_points), money(jira.story_points / 2), money(jira.story_points * 4), "Open Bookings →"])
        master.cell(mrow, 6).hyperlink = f"#'{sheet_name}'!A1"
        master.cell(mrow, 6).font = Font(name="Aptos", color=BLUE, underline="single", bold=True)
        ws = wb.create_sheet(sheet_name)
        ws.merge_cells("A1:E1")
        ws["A1"] = f"{jira.code}  |  BCS BOOKING HISTORY"
        ws.merge_cells("A2:E2")
        ws["A2"] = excel_text(f"Sprint: {jira.sprint}   •   Story Points: {money(jira.story_points):g}   •   Total Hours: {money(jira.story_points * 4):g}")
        ws.append([])
        ws.append(["Total Hours", "Date", "Booking Hours for BCS", "Remark for BCS", "Remaining Hours"])
        remaining = jira.story_points * 4
        ordered = sorted_bookings(jira)
        if ordered:
            for booking in ordered:
                remaining -= booking.hours
                ws.append([money(jira.story_points * 4), booking.book_date, money(booking.hours), excel_text(booking.remark), money(remaining)])
                ws.cell(ws.max_row, 2).number_format = "dd mmm yyyy"
        style_sheet(ws, ["Total Hours", "Date", "Booking Hours for BCS", "Remark for BCS", "Remaining Hours"], [18, 20, 26, 42, 23])
        ws["A1"].hyperlink = "#'Master Tracker'!A1"
        ws["A1"].font = Font(name="Aptos Display", size=19, color=BLUE, bold=True)
        for row in range(5, ws.max_row + 1):
            ws.cell(row, 5).font = Font(name="Aptos", color="C2410C" if ws.cell(row, 5).value < 0 else GREEN, bold=True)
    style_sheet(master, headers, [22, 20, 17, 15, 19, 22])
    for row in range(5, master.max_row + 1):
        master.cell(row, 6).font = Font(name="Aptos", color=BLUE, underline="single", bold=True)
        for col in (3, 4, 5):
            master.cell(row, col).number_format = "0.##"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@app.get("/export/excel")
def excel_export():
    jiras = db.session.execute(select(Jira).options(selectinload(Jira.bookings)).order_by(Jira.sprint, Jira.code)).scalars().all()
    return send_file(export_workbook(jiras), as_attachment=True, download_name=f"JIRA_BCS_Tracker_{date.today().isoformat()}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/export/excel/jira/<jira_id>")
def excel_jira_export(jira_id):
    jira = db.session.execute(select(Jira).where(Jira.id == jira_id).options(selectinload(Jira.bookings))).scalar_one_or_none()
    if not jira:
        abort(404)
    return send_file(export_workbook([jira]), as_attachment=True, download_name=f"{jira.code}_BCS_Bookings.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


with app.app_context():
    db.create_all()


if __name__ == "__main__":
    app.run(debug=not bool(os.environ.get("RENDER")), port=int(os.environ.get("PORT", 5000)))
